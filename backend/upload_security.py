from __future__ import annotations

import json
import multiprocessing
import os
import re
import shutil
import socket
import tempfile
import time
import unicodedata
import urllib.request
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path, PurePosixPath
from queue import Empty
from typing import Callable, Literal


UploadKind = Literal['pdf', 'docx', 'text']

ALLOWED_TEXT_EXTENSIONS = {'.txt', '.md', '.csv'}
ALLOWED_EXTENSIONS = {'.pdf', '.docx', *ALLOWED_TEXT_EXTENSIONS}
TEXT_CONTROL_WHITELIST = {9, 10, 12, 13}
NESTED_ARCHIVE_EXTENSIONS = {
    '.zip', '.7z', '.rar', '.tar', '.gz', '.bz2', '.xz', '.tgz', '.tbz', '.txz', '.iso'
}
EXECUTABLE_EXTENSIONS = {
    '.exe', '.dll', '.scr', '.bat', '.cmd', '.ps1', '.js', '.jse', '.vbs', '.vbe', '.jar', '.msi', '.com'
}
MACRO_MARKERS = {
    'word/vbaproject.bin', 'vbaProject.bin', 'word/activeprinter.bin'
}
PDF_ACTIVE_MARKERS = (
    b'/EmbeddedFile', b'/Filespec', b'/JavaScript', b'/JS', b'/OpenAction', b'/Launch', b'/AA'
)


@dataclass(frozen=True)
class UploadParserConfig:
    max_upload_bytes: int = 10 * 1024 * 1024
    max_filename_length: int = 180
    max_extracted_chars: int = 50_000
    max_pdf_pages: int = 30
    max_docx_entries: int = 100
    max_docx_total_uncompressed_bytes: int = 5 * 1024 * 1024
    max_docx_single_entry_uncompressed_bytes: int = 2 * 1024 * 1024
    max_docx_compression_ratio: float = 100.0
    parser_timeout_sec: float = 20.0
    max_parser_result_bytes: int = 100 * 1024


@dataclass(frozen=True)
class UploadParseResult:
    status: str
    message: str
    text: str = ''
    detected_type: str | None = None
    parser_invoked: bool = False
    temp_workspace_removed: bool = True
    child_pid: int | None = None

    @property
    def ok(self) -> bool:
        return self.status == 'ok'


DEFAULT_UPLOAD_PARSER_CONFIG = UploadParserConfig()

SAFE_MESSAGES = {
    'ok': 'Dokumen berhasil diproses secara aman.',
    'unsupported_type': 'Format dokumen tidak didukung. Unggah PDF, DOCX, atau teks UTF-8.',
    'type_mismatch': 'Jenis file tidak sesuai dengan isi dokumen.',
    'empty_file': 'Dokumen kosong dan tidak dapat dianalisis.',
    'file_too_large': 'Ukuran file melebihi batas upload.',
    'filename_too_long': 'Nama file terlalu panjang.',
    'pdf_page_limit_exceeded': 'Jumlah halaman PDF melebihi batas aman.',
    'docx_archive_limit_exceeded': 'Struktur DOCX melebihi batas arsip aman.',
    'docx_zip_bomb_suspected': 'DOCX ditolak karena pola kompresi berisiko.',
    'docx_path_traversal_rejected': 'DOCX ditolak karena nama anggota arsip tidak aman.',
    'docx_macro_content_rejected': 'DOCX ditolak karena mengandung konten makro atau aktif.',
    'docx_nested_archive_rejected': 'DOCX ditolak karena mengandung arsip bersarang.',
    'docx_external_relationship_rejected': 'DOCX ditolak karena mengandung relasi eksternal.',
    'text_decode_rejected': 'Dokumen teks harus berupa UTF-8 bersih tanpa profil biner.',
    'parser_timeout': 'Parser dokumen melewati batas waktu aman dan dihentikan.',
    'parser_crashed': 'Parser dokumen gagal secara aman.',
    'extracted_text_too_large': 'Teks hasil ekstraksi melebihi batas aman.',
    'photo_analysis_unsupported': 'Analisis foto klinis tidak tersedia pada mode aman ini.',
    'pdf_active_content_rejected': 'PDF ditolak karena mengandung konten aktif atau lampiran.',
}


def config_from_app(app_config) -> UploadParserConfig:
    return UploadParserConfig(
        max_upload_bytes=int(getattr(app_config, 'upload_max_bytes', DEFAULT_UPLOAD_PARSER_CONFIG.max_upload_bytes)),
        max_filename_length=int(getattr(app_config, 'upload_max_filename_length', DEFAULT_UPLOAD_PARSER_CONFIG.max_filename_length)),
        max_extracted_chars=int(getattr(app_config, 'upload_max_extracted_chars', DEFAULT_UPLOAD_PARSER_CONFIG.max_extracted_chars)),
        max_pdf_pages=int(getattr(app_config, 'upload_max_pdf_pages', DEFAULT_UPLOAD_PARSER_CONFIG.max_pdf_pages)),
        max_docx_entries=int(getattr(app_config, 'upload_max_docx_entries', DEFAULT_UPLOAD_PARSER_CONFIG.max_docx_entries)),
        max_docx_total_uncompressed_bytes=int(getattr(app_config, 'upload_max_docx_total_uncompressed_bytes', DEFAULT_UPLOAD_PARSER_CONFIG.max_docx_total_uncompressed_bytes)),
        max_docx_single_entry_uncompressed_bytes=int(getattr(app_config, 'upload_max_docx_single_entry_uncompressed_bytes', DEFAULT_UPLOAD_PARSER_CONFIG.max_docx_single_entry_uncompressed_bytes)),
        max_docx_compression_ratio=float(getattr(app_config, 'upload_max_docx_compression_ratio', DEFAULT_UPLOAD_PARSER_CONFIG.max_docx_compression_ratio)),
        parser_timeout_sec=float(getattr(app_config, 'upload_parser_timeout_sec', DEFAULT_UPLOAD_PARSER_CONFIG.parser_timeout_sec)),
        max_parser_result_bytes=int(getattr(app_config, 'upload_max_parser_result_bytes', DEFAULT_UPLOAD_PARSER_CONFIG.max_parser_result_bytes)),
    )


def safe_message(status: str) -> str:
    return SAFE_MESSAGES.get(status, SAFE_MESSAGES['parser_crashed'])


def rejection(status: str, detected_type: str | None = None, parser_invoked: bool = False) -> UploadParseResult:
    return UploadParseResult(status=status, message=safe_message(status), detected_type=detected_type, parser_invoked=parser_invoked)


def _basename(filename: str | None) -> str:
    raw = (filename or '').replace('\\', '/').split('/')[-1]
    return raw.strip()


def _extension(filename: str | None) -> str:
    return Path(_basename(filename)).suffix.lower()


def _extension_matches(kind: UploadKind, filename: str | None) -> bool:
    ext = _extension(filename)
    if kind == 'pdf':
        return ext == '.pdf'
    if kind == 'docx':
        return ext == '.docx'
    if kind == 'text':
        return ext in ALLOWED_TEXT_EXTENSIONS
    return False


def _looks_like_clean_text(raw: bytes) -> bool:
    if not raw:
        return False
    if raw.startswith((b'MZ', b'\x7fELF', b'\xca\xfe\xba\xbe')):
        return False
    try:
        text = raw.decode('utf-8')
    except UnicodeDecodeError:
        return False
    if '\x00' in text:
        return False
    controls = [ord(ch) for ch in text if ord(ch) < 32 and ord(ch) not in TEXT_CONTROL_WHITELIST]
    return len(controls) <= max(1, len(text) // 200)


def sniff_upload_type(raw: bytes) -> tuple[str | None, str | None]:
    if not raw:
        return None, 'empty_file'
    if raw.startswith(b'%PDF-'):
        return 'pdf', None
    if zipfile.is_zipfile(_BytesView(raw)):
        if _has_docx_required_members(raw):
            return 'docx', None
        return 'zip', 'unsupported_type'
    if _looks_like_clean_text(raw):
        return 'text', None
    return None, 'unsupported_type'


class _BytesView:
    def __init__(self, raw: bytes) -> None:
        import io
        self._bio = io.BytesIO(raw)

    def __getattr__(self, name: str):
        return getattr(self._bio, name)


def _has_docx_required_members(raw: bytes) -> bool:
    try:
        with zipfile.ZipFile(_BytesView(raw)) as zf:
            names = {info.filename for info in zf.infolist()}
            return '[Content_Types].xml' in names and 'word/document.xml' in names
    except zipfile.BadZipFile:
        return False


def validate_docx_container(raw: bytes, cfg: UploadParserConfig) -> UploadParseResult:
    try:
        with zipfile.ZipFile(_BytesView(raw)) as zf:
            infos = zf.infolist()
            if len(infos) > cfg.max_docx_entries:
                return rejection('docx_archive_limit_exceeded', 'docx')
            names_seen: set[str] = set()
            total_uncompressed = 0
            names = {info.filename for info in infos}
            if '[Content_Types].xml' not in names or 'word/document.xml' not in names:
                return rejection('unsupported_type', 'docx')
            for info in infos:
                normalized = _normalize_zip_member(info.filename)
                if normalized is None or _is_zip_symlink(info) or _is_encrypted_zip_member(info):
                    return rejection('docx_path_traversal_rejected', 'docx')
                lower_name = normalized.lower()
                if lower_name in names_seen:
                    return rejection('docx_archive_limit_exceeded', 'docx')
                names_seen.add(lower_name)
                if lower_name in {marker.lower() for marker in MACRO_MARKERS} or lower_name.endswith('.bin'):
                    return rejection('docx_macro_content_rejected', 'docx')
                suffix = PurePosixPath(lower_name).suffix
                if suffix in NESTED_ARCHIVE_EXTENSIONS:
                    return rejection('docx_nested_archive_rejected', 'docx')
                if suffix in EXECUTABLE_EXTENSIONS:
                    return rejection('docx_macro_content_rejected', 'docx')
                if info.file_size > cfg.max_docx_single_entry_uncompressed_bytes:
                    return rejection('docx_archive_limit_exceeded', 'docx')
                total_uncompressed += info.file_size
                if total_uncompressed > cfg.max_docx_total_uncompressed_bytes:
                    return rejection('docx_archive_limit_exceeded', 'docx')
                if _compression_ratio_too_high(info, cfg):
                    return rejection('docx_zip_bomb_suspected', 'docx')
                if lower_name.endswith('.rels') and _rels_has_external_target(zf, info):
                    return rejection('docx_external_relationship_rejected', 'docx')
    except zipfile.BadZipFile:
        return rejection('unsupported_type', 'docx')
    except Exception:
        return rejection('parser_crashed', 'docx')
    return UploadParseResult(status='ok', message=safe_message('ok'), detected_type='docx')


def _normalize_zip_member(name: str) -> str | None:
    if not name or '\x00' in name or '\\' in name:
        return None
    normalized_name = unicodedata.normalize('NFKC', name)
    if normalized_name != name:
        return None
    if name.startswith('/') or re.match(r'^[A-Za-z]:', name):
        return None
    path = PurePosixPath(name)
    if any(part in {'..', ''} for part in path.parts):
        return None
    return path.as_posix()


def _is_zip_symlink(info: zipfile.ZipInfo) -> bool:
    mode = (info.external_attr >> 16) & 0o170000
    return mode == 0o120000

def _is_encrypted_zip_member(info: zipfile.ZipInfo) -> bool:
    return bool(info.flag_bits & 0x1)


def _compression_ratio_too_high(info: zipfile.ZipInfo, cfg: UploadParserConfig) -> bool:
    if info.file_size == 0:
        return False
    if info.compress_size == 0:
        return True
    return (info.file_size / max(info.compress_size, 1)) > cfg.max_docx_compression_ratio


def _rels_has_external_target(zf: zipfile.ZipFile, info: zipfile.ZipInfo) -> bool:
    if info.file_size > 256 * 1024:
        return True
    try:
        body = zf.read(info.filename).decode('utf-8', 'ignore').lower()
    except Exception:
        return True
    compact = body.replace(chr(34), '').replace(chr(39), '')
    return 'targetmode=external' in compact


def parse_document_bytes(
    raw: bytes,
    filename: str | None,
    content_type: str | None = None,
    config: UploadParserConfig = DEFAULT_UPLOAD_PARSER_CONFIG,
    worker: Callable | None = None,
) -> UploadParseResult:
    del content_type
    if filename is not None and len(filename) > config.max_filename_length:
        return rejection('filename_too_long')
    if len(raw) == 0:
        return rejection('empty_file')
    if len(raw) > config.max_upload_bytes:
        return rejection('file_too_large')
    detected, sniff_error = sniff_upload_type(raw)
    if sniff_error and detected != 'zip':
        return rejection(sniff_error, detected)
    if detected == 'zip':
        return rejection('unsupported_type', 'zip')
    if detected not in {'pdf', 'docx', 'text'}:
        return rejection('unsupported_type')
    if not _extension_matches(detected, filename):
        return rejection('type_mismatch', detected)
    if detected == 'pdf' and _pdf_has_policy_disallowed_markers(raw):
        return rejection('pdf_active_content_rejected', 'pdf')
    if detected == 'docx':
        validation = validate_docx_container(raw, config)
        if not validation.ok:
            return validation
    if detected == 'text' and len(raw) > config.max_extracted_chars * 4:
        return rejection('extracted_text_too_large', 'text')
    return _run_parser_process(raw, detected, config, worker or _parser_worker)


def _run_parser_process(raw: bytes, kind: UploadKind, cfg: UploadParserConfig, worker: Callable) -> UploadParseResult:
    tmp_dir = tempfile.mkdtemp(prefix='cdss_upload_parse_')
    path = Path(tmp_dir) / 'upload.bin'
    child_pid: int | None = None
    result_queue = None
    try:
        Path(tmp_dir).mkdir(parents=True, exist_ok=True)
        path.write_bytes(raw)
        ctx = multiprocessing.get_context('spawn')
        result_queue = ctx.Queue(maxsize=1)
        proc = ctx.Process(target=worker, args=(str(path), kind, asdict(cfg), result_queue))
        proc.start()
        child_pid = int(proc.pid or 0) or None
        proc.join(cfg.parser_timeout_sec)
        if proc.is_alive():
            proc.terminate()
            proc.join(1.0)
            if proc.is_alive() and hasattr(proc, 'kill'):
                proc.kill()
                proc.join(1.0)
            return UploadParseResult(
                status='parser_timeout',
                message=safe_message('parser_timeout'),
                detected_type=kind,
                parser_invoked=True,
                child_pid=child_pid,
            )
        try:
            payload = result_queue.get_nowait()
        except Empty:
            return UploadParseResult(
                status='parser_crashed',
                message=safe_message('parser_crashed'),
                detected_type=kind,
                parser_invoked=True,
                child_pid=child_pid,
            )
        if len(json.dumps(payload, ensure_ascii=False).encode('utf-8')) > cfg.max_parser_result_bytes:
            return UploadParseResult(
                status='extracted_text_too_large',
                message=safe_message('extracted_text_too_large'),
                detected_type=kind,
                parser_invoked=True,
                child_pid=child_pid,
            )
        return _result_from_payload(payload, kind, child_pid)
    finally:
        if result_queue is not None:
            try:
                result_queue.close()
                result_queue.join_thread()
            except Exception:
                pass
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _result_from_payload(payload: dict, kind: str, child_pid: int | None) -> UploadParseResult:
    status = str(payload.get('status') or 'parser_crashed')
    text = str(payload.get('text') or '') if status == 'ok' else ''
    return UploadParseResult(
        status=status if status in SAFE_MESSAGES else 'parser_crashed',
        message=safe_message(status),
        text=text,
        detected_type=str(payload.get('detected_type') or kind),
        parser_invoked=True,
        child_pid=child_pid,
    )


def _parser_worker(path: str, kind: UploadKind, cfg_raw: dict, result_queue) -> None:
    cfg = UploadParserConfig(**cfg_raw)
    _install_parser_runtime_guards()
    try:
        if kind == 'pdf':
            result = _parse_pdf(path, cfg)
        elif kind == 'docx':
            result = _parse_docx(path, cfg)
        elif kind == 'text':
            result = _parse_text(path, cfg)
        else:
            result = {'status': 'unsupported_type', 'detected_type': kind}
    except Exception:
        result = {'status': 'parser_crashed', 'detected_type': kind}
    _queue_result(result_queue, result)


def _queue_result(result_queue, payload: dict) -> None:
    try:
        result_queue.put(payload, block=False)
    except Exception:
        pass


class ParserRuntimeBlocked(RuntimeError):
    pass


def _install_parser_runtime_guards() -> None:
    def blocked(*_args, **_kwargs):
        raise ParserRuntimeBlocked('parser outbound and command execution disabled')

    socket.socket.connect = blocked
    socket.create_connection = blocked
    socket.getaddrinfo = blocked
    urllib.request.urlopen = blocked
    os.system = blocked
    try:
        import subprocess
        subprocess.Popen = blocked
        subprocess.run = blocked
        subprocess.call = blocked
        subprocess.check_call = blocked
        subprocess.check_output = blocked
    except Exception:
        pass
    try:
        import requests
        requests.sessions.Session.request = blocked
        requests.get = blocked
        requests.post = blocked
    except Exception:
        pass
    try:
        import httpx
        httpx.Client.request = blocked
        httpx.AsyncClient.request = blocked
        httpx.get = blocked
        httpx.post = blocked
    except Exception:
        pass

def _pdf_has_policy_disallowed_markers(raw: bytes) -> bool:
    sample = raw[:2 * 1024 * 1024]
    return any(marker in sample for marker in PDF_ACTIVE_MARKERS)


def _parse_pdf(path: str, cfg: UploadParserConfig) -> dict:
    from PyPDF2 import PdfReader

    with open(path, 'rb') as fh:
        reader = PdfReader(fh, strict=False)
        if getattr(reader, 'is_encrypted', False):
            return {'status': 'unsupported_type', 'detected_type': 'pdf'}
        if len(reader.pages) > cfg.max_pdf_pages:
            return {'status': 'pdf_page_limit_exceeded', 'detected_type': 'pdf'}
        text = '\n'.join((page.extract_text() or '') for page in reader.pages)
    return _clean_text_payload(text, 'pdf', cfg)


def _parse_docx(path: str, cfg: UploadParserConfig) -> dict:
    raw = Path(path).read_bytes()
    validation = validate_docx_container(raw, cfg)
    if not validation.ok:
        return {'status': validation.status, 'detected_type': 'docx'}
    import docx

    document = docx.Document(path)
    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text for cell in row.cells))
    return _clean_text_payload('\n'.join(parts), 'docx', cfg)


def _parse_text(path: str, cfg: UploadParserConfig) -> dict:
    raw = Path(path).read_bytes()
    if not _looks_like_clean_text(raw):
        return {'status': 'text_decode_rejected', 'detected_type': 'text'}
    try:
        text = raw.decode('utf-8')
    except UnicodeDecodeError:
        return {'status': 'text_decode_rejected', 'detected_type': 'text'}
    return _clean_text_payload(text, 'text', cfg)


def _clean_text_payload(text: str, kind: str, cfg: UploadParserConfig) -> dict:
    cleaned = _sanitize_extracted_text(text)
    if len(cleaned) > cfg.max_extracted_chars:
        return {'status': 'extracted_text_too_large', 'detected_type': kind}
    return {'status': 'ok', 'detected_type': kind, 'text': cleaned}


def _sanitize_extracted_text(text: str) -> str:
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\x00', '')
    chars: list[str] = []
    for ch in text:
        code = ord(ch)
        if code >= 32 or code in TEXT_CONTROL_WHITELIST:
            chars.append(ch)
    normalized = ''.join(chars)
    normalized = re.sub(r'\n{4,}', '\n\n\n', normalized)
    return normalized.strip()


def wait_until_process_exits(pid: int | None, timeout_sec: float = 2.0) -> bool:
    if not pid:
        return True
    deadline = time.time() + timeout_sec
    while time.time() < deadline:
        if not _pid_alive(pid):
            return True
        time.sleep(0.05)
    return not _pid_alive(pid)


def _pid_alive(pid: int) -> bool:
    if pid <= 0:
        return False
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False
